from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor


def heading(document, text, level=1):
    document.add_heading(text, level=level)


def bullet(document, text, level=0):
    style = 'List Bullet' if level == 0 else 'List Bullet 2'
    document.add_paragraph(text, style=style)


def numbered(document, text):
    document.add_paragraph(text, style='List Number')


def make_doc(path: Path):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(10.5)
    for name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        styles[name].font.name = 'Aptos Display'
    styles['Title'].font.color.rgb = RGBColor(19, 124, 80)
    styles['Heading 1'].font.color.rgb = RGBColor(19, 124, 80)
    styles['Heading 2'].font.color.rgb = RGBColor(18, 32, 28)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Panduan Fungsional FACEDETECTION')
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(19, 124, 80)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run('Absensi wajah, GPS, cuti, lembur, dan payslip').italic = True
    document.add_paragraph('Versi: 1.0 | Untuk pengguna aplikasi Facework dan admin HR')

    heading(document, '1. Gambaran Umum', 1)
    document.add_paragraph(
        'FACEDETECTION adalah aplikasi presensi dan layanan HR yang menggunakan verifikasi wajah, '
        'lokasi GPS, dan integrasi Odoo. Odoo tetap menjadi sumber data utama untuk attendance, '
        'time off, overtime, dan payroll.'
    )
    bullet(document, 'Frontend pengguna: https://hris.kanjabung.com')
    bullet(document, 'Backend API: https://hris.kanjabung.web.id')
    bullet(document, 'Odoo HR: https://siskab.kanjabung.com')

    heading(document, '2. Persiapan Pengguna', 1)
    for item in [
        'Gunakan akun Odoo yang aktif dan terhubung dengan data karyawan.',
        'Gunakan browser modern seperti Chrome, Edge, atau Safari versi terbaru.',
        'Izinkan kamera ketika diminta oleh browser.',
        'Izinkan lokasi/GPS ketika diminta oleh browser.',
        'Pastikan jaringan internet stabil saat melakukan presensi.',
    ]:
        bullet(document, item)

    heading(document, '3. Login', 1)
    numbered(document, 'Buka https://hris.kanjabung.com.')
    numbered(document, 'Masukkan username dan password Odoo.')
    numbered(document, 'Pastikan alamat database Odoo sesuai konfigurasi perusahaan.')
    numbered(document, 'Tekan Login.')
    document.add_paragraph('Setelah berhasil, dashboard menampilkan identitas pengguna, status presensi, dan akses cepat ke cuti serta lembur.')

    heading(document, '4. Pendaftaran Wajah', 1)
    document.add_paragraph('Pendaftaran wajah dilakukan satu kali atau diulang apabila wajah berubah secara signifikan.')
    for item in [
        'Buka menu Wajah.',
        'Tekan mulai pendaftaran.',
        'Izinkan kamera.',
        'Posisikan wajah di dalam area panduan.',
        'Ambil beberapa sampel dari arah depan, kiri, kanan, atas, dan bawah.',
        'Pastikan pencahayaan cukup, wajah tidak tertutup, dan gambar tidak buram.',
        'Tunggu sampai status pendaftaran wajah aktif.',
    ]:
        numbered(document, item)

    heading(document, '5. Presensi Check-in dan Check-out', 1)
    heading(document, '5.1 Check-in', 2)
    for item in [
        'Buka menu Presensi.',
        'Pilih Check-in.',
        'Pastikan wajah berada di area kamera.',
        'Pastikan GPS aktif dan berada di lokasi kerja yang diizinkan.',
        'Tunggu proses verifikasi wajah dan GPS selesai.',
        'Periksa pesan hasil presensi dan status sinkronisasi Odoo.',
    ]:
        numbered(document, item)
    heading(document, '5.2 Check-out', 2)
    for item in [
        'Buka menu Presensi setelah selesai bekerja.',
        'Pilih Check-out.',
        'Ulangi proses verifikasi wajah dan lokasi.',
        'Pastikan hasil menunjukkan presensi berhasil.',
    ]:
        numbered(document, item)
    document.add_paragraph('Sistem mencegah check-in ganda dan mencatat event secara idempotent agar tidak terjadi data attendance duplikat.')

    heading(document, '6. Riwayat Presensi', 1)
    for item in [
        'Buka menu Riwayat.',
        'Periksa tanggal, waktu, jenis aktivitas, status, tingkat kemiripan wajah, dan informasi GPS.',
        'Jika sinkronisasi tertunda, status akan diproses ulang oleh worker integrasi.',
    ]:
        bullet(document, item)

    heading(document, '7. Pengajuan Cuti / Time Off', 1)
    for item in [
        'Buka menu Cuti.',
        'Pilih jenis cuti.',
        'Isi tanggal mulai dan tanggal selesai.',
        'Isi keterangan pengajuan.',
        'Tekan Kirim Pengajuan.',
        'Pantau status pengajuan pada daftar cuti. Persetujuan mengikuti alur Odoo.',
    ]:
        numbered(document, item)

    heading(document, '8. Pengajuan Lembur / Overtime', 1)
    for item in [
        'Buka menu Lembur.',
        'Pilih tanggal lembur.',
        'Isi durasi dalam jam.',
        'Isi keterangan pekerjaan lembur.',
        'Tekan Kirim Pengajuan.',
        'Pantau status pada daftar lembur.',
    ]:
        numbered(document, item)

    heading(document, '9. Payslip / Slip Gaji', 1)
    for item in [
        'Buka menu Gaji.',
        'Pilih periode payslip yang tersedia.',
        'Tekan lihat atau download PDF.',
        'Jika payslip tidak tersedia, hubungi admin payroll.',
    ]:
        numbered(document, item)

    heading(document, '10. Troubleshooting Pengguna', 1)
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ('Masalah', 'Solusi')):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    issues = [
        ('Kamera tidak muncul', 'Periksa izin kamera browser, tutup aplikasi lain yang menggunakan kamera, lalu muat ulang halaman.'),
        ('GPS ditolak', 'Aktifkan Location Services dan izinkan lokasi untuk hris.kanjabung.com.'),
        ('Wajah tidak cocok', 'Gunakan pencahayaan cukup, lepaskan masker/kacamata gelap, dan lihat kamera secara langsung.'),
        ('Di luar radius lokasi', 'Pastikan berada di area kerja yang dikonfigurasi HR dan GPS memiliki akurasi baik.'),
        ('Data belum masuk Odoo', 'Jangan mengulang berkali-kali. Periksa Riwayat atau hubungi admin untuk pengecekan antrean sinkronisasi.'),
        ('Login gagal', 'Pastikan username, password, dan database benar; bila tetap gagal hubungi admin Odoo.'),
    ]
    for problem, solution in issues:
        cells = table.add_row().cells
        cells[0].text = problem
        cells[1].text = solution

    heading(document, '11. Panduan Admin HR', 1)
    for item in [
        'Pastikan employee Odoo memiliki user, barcode/identification ID, dan data perusahaan yang benar.',
        'Atur lokasi attendance dan radius pada konfigurasi Face Attendance.',
        'Pastikan modul hr_attendance, hr_holidays, hr_payroll_community, grt_face_attendance_bridge, dan grt_external_api aktif.',
        'Gunakan user integrasi khusus untuk API dan jangan membagikan client secret.',
        'Periksa log api-facedetection dan Odoo jika integrasi mengalami kegagalan.',
    ]:
        bullet(document, item)

    heading(document, '12. Alur Integrasi Teknis', 1)
    document.add_paragraph('Alur normal aplikasi adalah:')
    for item in [
        'User login melalui frontend.',
        'Frontend mengirim permintaan ke FastAPI.',
        'FastAPI memverifikasi wajah dan GPS.',
        'FastAPI meminta JWT ke Odoo External API Gateway.',
        'FastAPI mengirim event attendance atau permintaan HR ke Odoo.',
        'Odoo memvalidasi scope, user, employee, dan aturan HR.',
        'Hasil dikembalikan ke frontend dan dicatat pada sistem terkait.',
    ]:
        numbered(document, item)

    heading(document, '13. Keamanan dan Privasi', 1)
    bullet(document, 'Jangan membagikan password Odoo atau client secret API.')
    bullet(document, 'Gunakan perangkat pribadi yang aman dan browser yang diperbarui.')
    bullet(document, 'Data wajah dan presensi hanya digunakan untuk kebutuhan operasional HR sesuai kebijakan perusahaan.')
    bullet(document, 'Logout setelah selesai menggunakan aplikasi pada perangkat bersama.')

    document.add_paragraph()
    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('FACEDETECTION / Facework — Panduan Fungsional').italic = True
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


if __name__ == '__main__':
    import sys
    make_doc(Path(sys.argv[1]))
